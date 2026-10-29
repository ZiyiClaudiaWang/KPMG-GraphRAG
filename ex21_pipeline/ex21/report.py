"""Human-readable Excel report."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

F = "Arial"
HEAD = PatternFill("solid", fgColor="1F3864")
BAND = PatternFill("solid", fgColor="EEF2F8")
NEW = PatternFill("solid", fgColor="E2F0D9")
GONE = PatternFill("solid", fgColor="FBE5D6")
RENAME = PatternFill("solid", fgColor="FFF2CC")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))


def _table(ws, df: pd.DataFrame, start_row: int = 1, widths: dict | None = None):
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(start_row, j, col)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, rec in enumerate(df.itertuples(index=False), start_row + 1):
        for j, v in enumerate(rec, 1):
            if isinstance(v, float) and pd.isna(v):
                v = None
            c = ws.cell(i, j, v)
            c.font = Font(name=F)
            c.border = THIN
            c.alignment = Alignment(vertical="top", wrap_text=False)
    last = start_row + len(df)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(df.columns))}{max(last, start_row)}"
    ws.freeze_panes = ws.cell(start_row + 1, 2)
    for j, col in enumerate(df.columns, 1):
        w = (widths or {}).get(col)
        if w is None:
            sample = [len(str(col))] + [len(str(x)) for x in df.iloc[:, j - 1].head(200)]
            w = min(max(sample) + 2, 55)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[start_row].height = 30


def write_workbook(path: Path, df, tl, ev, qa):
    wb = Workbook()

    # 1. README ---------------------------------------------------------------
    ws = wb.active
    ws.title = "README"
    parents = ", ".join(f"{t} ({n})" for t, n in df[["parent_ticker", "parent_name"]].drop_duplicates().itertuples(index=False))
    lines = [
        ("Exhibit 21 Parent–Subsidiary Report", True, 14),
        ("", False, 10),
        ("Companies covered: " + parents, False, 10),
        ("Source: Exhibit 21 (\"Subsidiaries of the Registrant\") attached to each company's Form 10-K on SEC EDGAR. Every row keeps its accession number and source link.", False, 10),
        ("", False, 10),
        ("Sheets", True, 11),
        ("Summary — subsidiaries listed and countries covered, per company per fiscal year, plus additions / removals / renames vs the prior year.", False, 10),
        ("Parent → Subsidiaries — the latest list for each company. Start here.", False, 10),
        ("Timeline — one row per subsidiary; ✓ shows which years it was listed; status = listed every year / new / no longer listed / name changed.", False, 10),
        ("Changes — year-over-year events: Added, Removed, Renamed, Jurisdiction changed.", False, 10),
        ("All Records — every parsed row (one row = one subsidiary in one 10-K) with full provenance.", False, 10),
        ("QA — parsing checks per filing: rows parsed, duplicates, missing jurisdictions, files that need review.", False, 10),
        ("", False, 10),
        ("How to read it", True, 11),
        ("1. Exhibit 21 lists significant subsidiaries only (Reg. S-K Item 601(b)(21)); companies may omit subsidiaries that together are not significant. It is not a complete list.", False, 10),
        ("2. When the filing is a flat list it does not say whether a subsidiary is held directly or through another subsidiary, so the relationship reads 'direct/indirect not stated'. If the filing uses indentation, the pipeline records the immediate parent.", False, 10),
        ("3. 'No longer listed' only means the entity is absent from that year's Exhibit 21: it may have merged, been sold, been renamed, or simply dropped below the significance threshold. Check the 10-K text before drawing conclusions.", False, 10),
        ("4. Cross-year matching: identical names are the same entity; names that match once legal suffixes (Inc., Limited, Unlimited Company, ...) are removed, and are unambiguous that year, are treated as a rename.", False, 10),
    ]
    for i, (t, b, s) in enumerate(lines, 1):
        c = ws.cell(i, 1, t)
        c.font = Font(name=F, bold=b, size=s)
    ws.column_dimensions["A"].width = 130

    # 2. Summary (formulas over All Records) ------------------------------------
    ws = wb.create_sheet("Summary")
    pairs = df[["parent_ticker", "parent_name", "fiscal_year"]].drop_duplicates().sort_values(["parent_ticker", "fiscal_year"])
    head = ["Ticker", "Parent company", "Fiscal year", "Subsidiaries listed", "Countries",
            "Added vs prior year", "Removed vs prior year", "Renamed"]
    for j, h in enumerate(head, 1):
        c = ws.cell(1, j, h)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    n_all = len(df) + 1
    n_ev = len(ev) + 1
    for i, (t, name, fy) in enumerate(pairs.itertuples(index=False), 2):
        ws.cell(i, 1, t); ws.cell(i, 2, name); ws.cell(i, 3, int(fy))
        ws.cell(i, 4, f"=COUNTIFS('All Records'!$A$2:$A${n_all},A{i},'All Records'!$D$2:$D${n_all},C{i})")
        n_c = df[(df.parent_ticker == t) & (df.fiscal_year == fy)]["country"].nunique()
        c = ws.cell(i, 5, n_c)  # distinct count computed in Python (no reliable COUNTIFS equivalent)
        for col, kind in ((6, "Added"), (7, "Removed"), (8, "Renamed")):
            ws.cell(i, col, f"=COUNTIFS(Changes!$A$2:$A${n_ev},A{i},Changes!$B$2:$B${n_ev},C{i},Changes!$C$2:$C${n_ev},\"{kind}\")")
        for j in range(1, 9):
            ws.cell(i, j).font = Font(name=F)
            ws.cell(i, j).border = THIN
            if i % 2 == 0:
                ws.cell(i, j).fill = BAND
    note_row = len(pairs) + 3
    ws.cell(note_row, 1, "Countries = distinct countries in that year's list (computed by the pipeline). Other counts are formulas over the All Records and Changes sheets.").font = Font(name=F, italic=True, size=9)
    for col, w in zip("ABCDEFGH", (9, 26, 11, 13, 11, 13, 14, 10)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    # 3. Parent -> Subsidiaries (latest year per parent) ------------------------
    ws = wb.create_sheet("Parent → Subsidiaries")
    latest = df[df["fiscal_year"] == df.groupby("parent_ticker")["fiscal_year"].transform("max")]
    st = tl.set_index("entity_id")["status"]
    view = pd.DataFrame({
        "Parent company": latest["parent_name"] + " (" + latest["parent_ticker"] + ")",
        "Fiscal year": "FY" + latest["fiscal_year"].astype(str),
        "Subsidiary": latest["subsidiary_name"],
        "Immediate parent (per filing)": latest["immediate_parent"],
        "Country": latest["country"],
        "State / region": latest["state_or_region"],
        "Jurisdiction (as filed)": latest["jurisdiction_raw"],
        "3-year status": latest["entity_id"].map(st),
        "Source (accession)": latest["accession"],
    }).sort_values(["Parent company", "Country", "Subsidiary"])
    _table(ws, view)
    prev = None
    for r in range(2, ws.max_row + 1):
        cur = ws.cell(r, 1).value
        if cur != prev:
            for j in range(1, view.shape[1] + 1):
                ws.cell(r, j).border = Border(top=Side(style="medium", color="1F3864"))
            ws.cell(r, 1).font = Font(name=F, bold=True)
        else:
            ws.cell(r, 1).font = Font(name=F, color="808080")
        prev = cur
        s = ws.cell(r, 8).value or ""
        fill = NEW if s.startswith("New") else RENAME if "name changed" in s else None
        if fill:
            ws.cell(r, 8).fill = fill

    # 4. Timeline ---------------------------------------------------------------
    ws = wb.create_sheet("Timeline")
    fy_cols = sorted([c for c in tl.columns if c.startswith("FY")])
    tview = tl[["parent_ticker", "subsidiary_name (latest)", "earlier names", "country",
                "jurisdiction (as filed)"] + fy_cols + ["status", "entity_id"]]
    tview = tview.rename(columns={"parent_ticker": "Parent"})
    _table(ws, tview, widths={c: 8 for c in fy_cols})
    for r in range(2, ws.max_row + 1):
        for j, col in enumerate(tview.columns, 1):
            if col in fy_cols:
                ws.cell(r, j).alignment = Alignment(horizontal="center")
        s = ws.cell(r, tview.columns.get_loc("status") + 1)
        v = s.value or ""
        s.fill = NEW if v.startswith("New") else GONE if v.startswith("No longer") else RENAME if "name changed" in v else PatternFill()
    ws.cell(ws.max_row + 2, 1, "Blank FY cell = that parent did not file for that year in this dataset, or the subsidiary was not listed.").font = Font(name=F, italic=True, size=9)

    # 5. Changes ----------------------------------------------------------------
    ws = wb.create_sheet("Changes")
    _table(ws, ev, widths={"detail": 50})
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 3).value
        ws.cell(r, 3).fill = {"Added": NEW, "Removed": GONE, "Renamed": RENAME}.get(v, PatternFill())

    # 6. All Records --------------------------------------------------------------
    ws = wb.create_sheet("All Records")
    cols = ["parent_ticker", "parent_name", "parent_cik", "fiscal_year", "period_end", "accession",
            "subsidiary_name", "immediate_parent", "relationship", "level", "jurisdiction_raw",
            "country", "state_or_region", "ownership_pct", "entity_id", "row_order", "flags",
            "name_change_note", "source_url"]
    _table(ws, df[cols].sort_values(["parent_ticker", "fiscal_year", "row_order"]),
           widths={"source_url": 40, "relationship": 30})

    # 7. QA -------------------------------------------------------------------------
    ws = wb.create_sheet("QA")
    _table(ws, qa)
    wb.save(path)
